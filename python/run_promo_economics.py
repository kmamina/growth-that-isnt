import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config import ROOT

DB = ROOT / "data" / "zesto.duckdb"
pd.set_option("display.width", 160)

con = duckdb.connect(str(DB), read_only=True)

# 1. Promo economics, all valid orders
promo = con.execute("""
    SELECT COALESCE(p.promo_type, 'no_promo') AS promo_type,
           COUNT(*) AS orders,
           ROUND(SUM(o.order_value), 0) AS gmv,
           ROUND(SUM(o.discount_amount), 0) AS discount,
           ROUND(SUM(o.discount_amount) / SUM(o.order_value), 4) AS discount_rate,
           ROUND(AVG(o.cm), 2) AS cm_per_order
    FROM v_orders_q o LEFT JOIN clean_promotions p ON p.promo_id = o.promo_id
    GROUP BY 1 ORDER BY orders DESC""").fetchdf()
print("PROMO ECONOMICS (all valid orders)\n", promo.to_string(index=False))

# 2. 90-day CM per new customer, grouped by first-order promo
cust = con.execute("""
    SELECT f.customer_id, f.q,
           CASE WHEN f.promo_type = 'first_order_flat' THEN 'first_order_flat'
                WHEN f.promo_type IS NULL THEN 'no_promo' ELSE 'other_promo' END AS grp,
           f.repeat_90d, SUM(o.cm) AS cm90, COUNT(*) AS orders90
    FROM v_first_orders f
    JOIN v_orders_q o ON o.customer_id = f.customer_id
     AND o.order_ts >= f.first_ts AND o.order_ts <= f.first_ts + INTERVAL 90 DAY
    GROUP BY 1, 2, 3, 4""").fetchdf()
con.close()

g = (cust[cust.q.isin(["Q-1", "Q0"])]
     .groupby(["q", "grp"])
     .agg(customers=("customer_id", "size"), repeat_90d=("repeat_90d", "mean"),
          cm90=("cm90", "mean"), orders90=("orders90", "mean")).round(3).reset_index())
print("\n90-DAY VALUE PER NEW CUSTOMER\n", g.to_string(index=False))

promo.to_csv(ROOT / "reports" / "promo_economics.csv", index=False)
g.to_csv(ROOT / "reports" / "group_90d_cm.csv", index=False)

# 3. Excel model with live formulas
q0 = g[g.q == "Q0"].set_index("grp").loc[["first_order_flat", "other_promo", "no_promo"]]
wb = Workbook()
BLUE = Font(color="0000FF", bold=True)
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")


def header(ws, row, labels):
    for i, t in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


# Sheet: Data_Groups (Q0 cohort, values from the cleaned data)
ws = wb.active
ws.title = "Data_Groups"
header(ws, 1, ["First-order group (Q0 cohort)", "Customers", "90-day repeat rate",
               "90-day CM per customer ($)", "Orders in 90 days"])
for r, (grp, row) in enumerate(q0.iterrows(), start=2):
    ws.cell(r, 1, grp)
    ws.cell(r, 2, int(row.customers))
    ws.cell(r, 3, float(row.repeat_90d)).number_format = "0.0%"
    ws.cell(r, 4, float(row.cm90)).number_format = "0.00"
    ws.cell(r, 5, float(row.orders90)).number_format = "0.00"
ws.column_dimensions["A"].width = 32
for col in "BCDE":
    ws.column_dimensions[col].width = 20

# Sheet: Promo_Economics
wp = wb.create_sheet("Promo_Economics")
header(wp, 1, list(promo.columns))
for r, row in enumerate(promo.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        wp.cell(r, c, v)
for col in "ABCDEF":
    wp.column_dimensions[col].width = 18

# Sheet: Scenario
sc = wb.create_sheet("Scenario", 0)
sc["A1"] = "Scenario: remove the first_order_flat promo (Q0 new-customer cohort, 90-day view)"
sc["A1"].font = Font(bold=True, size=13)
rows = [
    (3, "DATA (from Data_Groups)", None),
    (4, "Q0 new customers", "=SUM(Data_Groups!B2:B4)"),
    (5, "Share of new customers on first_order_flat", "=Data_Groups!B2/B4"),
    (6, "90-day CM per customer, flat-promo group ($)", "=Data_Groups!D2"),
    (7, "90-day CM per customer, all other new customers ($)",
        "=(Data_Groups!B3*Data_Groups!D3+Data_Groups!B4*Data_Groups!D4)/(Data_Groups!B3+Data_Groups!B4)"),
    (8, "90-day repeat rate, flat-promo group", "=Data_Groups!C2"),
    (9, "90-day repeat rate, all other new customers",
        "=(Data_Groups!B3*Data_Groups!C3+Data_Groups!B4*Data_Groups!C4)/(Data_Groups!B3+Data_Groups!B4)"),
    (11, "ASSUMPTIONS (blue cells are editable)", None),
    (12, "Share of flat-promo customers who never sign up without the promo", 0.20),
    (13, "Value of the ones who stay, as % of a normal non-flat customer", 0.80),
    (15, "RESULTS", None),
    (16, "Baseline 90-day CM, Q0 cohort ($)", "=B4*B5*B6+B4*(1-B5)*B7"),
    (17, "Scenario 90-day CM ($)", "=B4*(1-B5)*B7+B4*B5*(1-B12)*B7*B13"),
    (18, "Change in 90-day CM ($)", "=B17-B16"),
    (19, "New customers lost", "=B4*B5*B12"),
    (20, "Baseline 90-day repeat rate", "=(B4*B5*B8+B4*(1-B5)*B9)/B4"),
    (21, "Scenario 90-day repeat rate", "=(B4*(1-B5)*B9+B4*B5*(1-B12)*B9)/(B4-B19)"),
    (22, "Break-even share lost (capped at 100%)", "=MAX(0,MIN(1,1-B6/(B7*B13)))"),
]
for r, label, val in rows:
    sc.cell(r, 1, label)
    if val is None:
        sc.cell(r, 1).font = Font(bold=True)
        continue
    c = sc.cell(r, 2, val)
    if r in (12, 13):
        c.font = BLUE
        c.number_format = "0%"
    elif r in (5, 8, 9, 20, 21, 22):
        c.number_format = "0.0%"
    elif r in (4, 19):
        c.number_format = "#,##0"
    elif r in (16, 17, 18):
        c.number_format = "#,##0"
    else:
        c.number_format = "0.00"
sc["A24"] = ("Notes: the scenario assumes customers who stay behave like other new customers; "
             "a lower value in B13 is a haircut for their higher promo sensitivity. "
             "This is an assumption, not a measured effect. Lost customers also mean lost GMV and growth, "
             "which this CM-only view does not price.")
sc["A24"].alignment = Alignment(wrap_text=True, vertical="top")
sc.merge_cells("A24:C27")
sc.column_dimensions["A"].width = 66
sc.column_dimensions["B"].width = 18

# Sheet: Sensitivity (change in 90-day CM, $)
se = wb.create_sheet("Sensitivity")
se["A1"] = "Change in 90-day CM ($): share lost (rows) x value of those who stay (columns)"
se["A1"].font = Font(bold=True)
conv = [0.5, 0.7, 0.8, 0.9, 1.0]
lost = [0.0, 0.1, 0.2, 0.3, 0.4, 0.5, 0.75, 1.0]
se.cell(2, 1, "Share lost \\ value kept")
for j, v in enumerate(conv, start=2):
    c = se.cell(2, j, v)
    c.number_format = "0%"
    c.font = HEAD
    c.fill = FILL
for i, l in enumerate(lost, start=3):
    c = se.cell(i, 1, l)
    c.number_format = "0%"
    c.font = Font(bold=True)
    for j in range(2, 2 + len(conv)):
        col = se.cell(2, j).column_letter
        se.cell(i, j, f"=Scenario!$B$4*Scenario!$B$5*((1-$A{i})*Scenario!$B$7*{col}$2-Scenario!$B$6)"
                ).number_format = "#,##0"
se.column_dimensions["A"].width = 24
for col in "BCDEF":
    se.column_dimensions[col].width = 14

out = ROOT / "excel" / "promo_scenario_model.xlsx"
wb.save(out)
print("\nSaved", out)

# Python check of the default scenario, so you can verify Excel's numbers
n = q0.customers.sum()
f = q0.loc["first_order_flat", "customers"] / n
cf = q0.loc["first_order_flat", "cm90"]
oth = q0.loc[["other_promo", "no_promo"]]
c_o = (oth.customers * oth.cm90).sum() / oth.customers.sum()
L, k = 0.20, 0.80
base = n * f * cf + n * (1 - f) * c_o
scen = n * (1 - f) * c_o + n * f * (1 - L) * c_o * k
print(f"CHECK (L=20%, keep=80%): baseline ${base:,.0f} | scenario ${scen:,.0f} | change ${scen-base:,.0f}"
      f" | customers lost {n*f*L:,.0f}")